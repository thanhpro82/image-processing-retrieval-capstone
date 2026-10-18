from __future__ import annotations

import argparse
import hashlib
import json
import random
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from statistics import median
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from PIL import Image, ImageOps

SOURCES = [
    (1, r"jobs\batch_1_1_den_1000_PHAN_CONG.xlsx"),
    (2, r"jobs\batch_2_1001_den_2450_PHAN_CONG.xlsx"),
    (3, r"jobs\batch_3_2451_den_3950_PHAN_CONG.xlsx"),
    (4, r"jobs\batch_4_3951_den_5450_PHAN_CONG_v3_20261009.xlsx"),
    (5, r"jobs\batch005\batch_5_5451_den_6950_PHAN_CONG.xlsx"),
    (6, r"jobs\batch_6_6951_den_8450_PHAN_CONG.xlsx"),
    (7, r"jobs\batch 7\batch_7_8451_den_10200_PHAN_CONG.xlsx"),
    (8, r"jobs\batch 8\batch_8_10201_den_11400_PHAN_CONG.xlsx"),
]
MAX_DOWNLOAD_BYTES = 20 * 1024 * 1024


def sample_rows(workspace: Path, per_batch: int, seed: int) -> list[dict]:
    selected: list[dict] = []
    globally_seen: set[str] = set()
    for batch, rel_path in SOURCES:
        source = workspace / rel_path
        if not source.is_file():
            raise FileNotFoundError(source)
        wb = load_workbook(source, read_only=True, data_only=True)
        ws = wb["Ảnh đã thẩm định"]
        ws.reset_dimensions()
        records: list[dict] = []
        local_seen: set[str] = set()
        for row in ws.iter_rows(min_row=2, min_col=1, max_col=11, values_only=True):
            stt, name, entity_id, order, _representative, _score, description, url, _resolution, page_url, title = row
            if not isinstance(url, str) or not url.startswith(("http://", "https://")):
                continue
            if url in local_seen or url in globally_seen:
                continue
            local_seen.add(url)
            records.append({
                "batch": batch,
                "stt": stt,
                "entity_name": name,
                "entity_id": entity_id,
                "image_order": order,
                "description": description,
                "source_url": url,
                "source_page": page_url,
                "source_title": title,
            })
        wb.close()
        if len(records) < per_batch:
            raise RuntimeError(f"Batch {batch}: only {len(records)} unique URLs; need {per_batch}")
        rng = random.Random(seed + batch)
        chosen = rng.sample(records, per_batch)
        selected.extend(chosen)
        globally_seen.update(r["source_url"] for r in chosen)
        print(f"sampled batch {batch}: {len(chosen)} unique URLs", flush=True)
    return selected


def download_one(record: dict, output_dir: Path, max_side: int) -> dict:
    started = time.perf_counter()
    url = record["source_url"]
    result = {**record, "http_status": None, "content_type": None, "source_bytes": 0,
              "elapsed_sec": 0.0, "status": "error", "error": None}
    try:
        request = Request(url, headers={
            "User-Agent": "Mozilla/5.0 (compatible; CapstoneImagePilot/1.0)",
            "Accept": "image/avif,image/webp,image/apng,image/*,*/*;q=0.8",
        })
        with urlopen(request, timeout=20) as response:
            result["http_status"] = response.status
            result["content_type"] = response.headers.get("Content-Type", "").split(";")[0]
            raw = response.read(MAX_DOWNLOAD_BYTES + 1)
        result["source_bytes"] = len(raw)
        if len(raw) > MAX_DOWNLOAD_BYTES:
            raise ValueError("response exceeds 20 MiB pilot limit")
        source_sha = hashlib.sha256(raw).hexdigest()
        with Image.open(BytesIO(raw)) as image:
            image.verify()
        with Image.open(BytesIO(raw)) as image:
            image = ImageOps.exif_transpose(image)
            original_size = list(image.size)
            image.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
            if image.mode not in ("RGB", "RGBA"):
                image = image.convert("RGB")
            elif image.mode == "RGBA":
                background = Image.new("RGB", image.size, "white")
                background.paste(image, mask=image.getchannel("A"))
                image = background
            name_hash = hashlib.sha256(url.encode("utf-8")).hexdigest()[:12]
            stt = str(record.get("stt") or "row").replace("/", "_").replace("\\", "_")
            filename = f"b{record['batch']}_{stt}_{name_hash}.webp"
            path = output_dir / filename
            image.save(path, format="WEBP", quality=82, method=4)
        result.update({
            "status": "valid_image",
            "source_sha256": source_sha,
            "original_dimensions": original_size,
            "processed_dimensions": list(image.size),
            "image_path": str(path.name),
            "processed_bytes": path.stat().st_size,
        })
    except HTTPError as exc:
        result["http_status"] = exc.code
        result["error"] = f"HTTP {exc.code}"
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
    result["elapsed_sec"] = round(time.perf_counter() - started, 4)
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", type=Path, default=Path(r"C:\Drive_D\Crawl Data"))
    parser.add_argument("--output", type=Path, default=Path(r"C:\Drive_D\capstone-image-retrieval\data\pilot"))
    parser.add_argument("--per-batch", type=int, default=125)
    parser.add_argument("--max-side", type=int, default=768)
    parser.add_argument("--workers", type=int, default=8)
    parser.add_argument("--seed", type=int, default=20261017)
    args = parser.parse_args()
    image_dir = args.output / "images"
    image_dir.mkdir(parents=True, exist_ok=True)
    records = sample_rows(args.workspace, args.per_batch, args.seed)
    started = time.perf_counter()
    results = []
    with ThreadPoolExecutor(max_workers=args.workers) as executor:
        futures = [executor.submit(download_one, row, image_dir, args.max_side) for row in records]
        for i, future in enumerate(as_completed(futures), 1):
            result = future.result()
            results.append(result)
            if i % 100 == 0 or i == len(futures):
                print(f"processed {i}/{len(futures)}; valid={sum(r['status']=='valid_image' for r in results)}", flush=True)
    wall_sec = time.perf_counter() - started

    # Remove byte-identical copies while preserving each URL/entity in the manifest.
    canonical_by_sha: dict[str, str] = {}
    unique_images = 0
    for result in results:
        if result["status"] != "valid_image":
            continue
        sha = result["source_sha256"]
        if sha in canonical_by_sha:
            duplicate_path = image_dir / result["image_path"]
            duplicate_path.unlink(missing_ok=True)
            result["duplicate_of"] = canonical_by_sha[sha]
            result["image_path"] = canonical_by_sha[sha]
            result["processed_bytes"] = 0
        else:
            canonical_by_sha[sha] = result["image_path"]
            unique_images += 1

    results.sort(key=lambda r: (r["batch"], str(r.get("stt"))))
    manifest_path = args.output / "manifest.jsonl"
    with manifest_path.open("w", encoding="utf-8", newline="\n") as stream:
        for result in results:
            stream.write(json.dumps(result, ensure_ascii=False) + "\n")
    valid = [r for r in results if r["status"] == "valid_image"]
    status_counts: dict[str, int] = {}
    per_batch: dict[str, dict] = {}
    for batch in range(1, 9):
        subset = [r for r in results if r["batch"] == batch]
        valid_batch = [r for r in subset if r["status"] == "valid_image"]
        per_batch[str(batch)] = {
            "sampled": len(subset),
            "valid_images": len(valid_batch),
            "errors": len(subset) - len(valid_batch),
            "source_bytes_valid": sum(r["source_bytes"] for r in valid_batch),
            "processed_bytes_unique": sum(r.get("processed_bytes", 0) for r in valid_batch),
        }
    for r in results:
        status_counts[r["status"]] = status_counts.get(r["status"], 0) + 1
    source_bytes = sum(r["source_bytes"] for r in valid)
    output_bytes = sum(r.get("processed_bytes", 0) for r in valid)
    report = {
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "sample_method": f"{args.per_batch} random unique URLs per batch from the reviewed-image sheet; duplicates across selected URLs excluded",
        "seed": args.seed,
        "max_side_px": args.max_side,
        "webp_quality": 82,
        "workers": args.workers,
        "sampled_urls": len(results),
        "valid_images": len(valid),
        "valid_rate": round(len(valid) / len(results), 4) if results else 0,
        "unique_byte_images": unique_images,
        "status_counts": status_counts,
        "elapsed_download_and_process_sec": round(wall_sec, 2),
        "valid_images_per_sec": round(len(valid) / wall_sec, 3) if wall_sec else 0,
        "source_bytes_valid": source_bytes,
        "processed_bytes_unique": output_bytes,
        "mean_source_bytes_per_valid": round(source_bytes / len(valid)) if valid else 0,
        "mean_processed_bytes_per_unique_image": round(output_bytes / unique_images) if unique_images else 0,
        "estimated_80000_processed_bytes": round(output_bytes / unique_images * 80000) if unique_images else None,
        "estimated_80000_elapsed_hours_at_measured_rate": round(80000 / (len(valid) / wall_sec) / 3600, 2) if valid and wall_sec else None,
        "by_batch": per_batch,
        "limitations": ["URL liveness and decoding only; no entity relevance review", "Exact-byte deduplication only", "Sample is not a guarantee of full-corpus download rate"],
    }
    (args.output / "pilot_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print("REPORT", json.dumps(report, ensure_ascii=False), flush=True)
    print(f"MANIFEST {manifest_path}", flush=True)


if __name__ == "__main__":
    main()
